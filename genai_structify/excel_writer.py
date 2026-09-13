"""Writes extracted rows to a styled .xlsx file."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .extractor import Schema

HEADER_FILL = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def write_excel(rows: list[dict], schema: Schema, output_path: str | Path) -> Path:
    output_path = Path(output_path)
    wb = Workbook()
    ws = wb.active
    ws.title = schema.sheet_name[:31]  # Excel sheet name limit

    headers = [f.name for f in schema.fields]
    ws.append(headers)

    for col_idx, _ in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")

    for row in rows:
        ws.append([row.get(h) for h in headers])

    _auto_size_columns(ws, headers, rows)
    ws.freeze_panes = "A2"

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    return output_path


def _auto_size_columns(ws, headers: list[str], rows: list[dict]) -> None:
    for col_idx, header in enumerate(headers, start=1):
        max_len = len(header)
        for row in rows:
            value = row.get(header)
            if value is not None:
                max_len = max(max_len, len(str(value)))
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 4, 60)
